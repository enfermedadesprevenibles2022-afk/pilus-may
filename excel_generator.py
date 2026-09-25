from __future__ import annotations

from copy import copy
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

MAX_PERSONAS = 100
PERSONAS_PLANTILLA = 15
MAX_SINTOMAS = 18
MAX_ALIMENTOS = 16

SINTOMA_COLS = [chr(c) for c in range(ord("I"), ord("Z") + 1)]
ALIMENTO_COLS = [chr(c) for c in range(ord("G"), ord("V") + 1)]

PERIODOS = [
    "Alimentos ingeridos día de los síntomas",
    "Alimentos ingeridos un día antes",
    "Alimentos ingeridos dos días antes",
]


def _clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    return str(value).strip()


def _truthy(value) -> bool:
    if isinstance(value, bool):
        return value
    text = _clean_text(value).casefold()
    return text in {"x", "si", "sí", "s", "1", "true", "verdadero", "yes"}


def _parse_date(text: str):
    text = _clean_text(text)
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _safe_filename(value: str) -> str:
    value = _clean_text(value) or "sin_fecha"
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_") or "sin_fecha"


def validate_payload(payload: dict) -> list[str]:
    errors: list[str] = []
    symptoms = payload.get("symptoms", [])
    foods = payload.get("foods", [])
    persons = payload.get("persons", [])
    food_rows = payload.get("food_rows", [])

    if len(symptoms) > MAX_SINTOMAS:
        errors.append(f"La plantilla admite máximo {MAX_SINTOMAS} signos/síntomas.")
    if len(foods) > MAX_ALIMENTOS:
        errors.append(f"La plantilla admite máximo {MAX_ALIMENTOS} alimentos.")
    if len(persons) > MAX_PERSONAS:
        errors.append(f"La aplicación admite máximo {MAX_PERSONAS} personas.")

    valid_nos = {str(i) for i in range(1, MAX_PERSONAS + 1)}
    for row in food_rows:
        n = _clean_text(row.get("No."))
        # Pandas puede convertir números enteros a texto con .0.
        try:
            n = str(int(float(n))) if n else ""
        except (TypeError, ValueError):
            pass
        if n and n not in valid_nos:
            errors.append(f"Número de persona inválido en consumo: {n}.")
            break
    return errors


def _copy_cell_format(source, target) -> None:
    """Copia el formato de una celda sin copiar su contenido."""
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _copy_row_format(ws, source_row: int, target_row: int, max_col: int) -> None:
    src_dim = ws.row_dimensions[source_row]
    dst_dim = ws.row_dimensions[target_row]
    dst_dim.height = src_dim.height
    dst_dim.hidden = src_dim.hidden
    dst_dim.outlineLevel = src_dim.outlineLevel
    dst_dim.collapsed = src_dim.collapsed

    for col in range(1, max_col + 1):
        _copy_cell_format(ws.cell(source_row, col), ws.cell(target_row, col))


def _expand_people_sheet(ws, person_count: int) -> int:
    """Amplía Hoja1 y devuelve la fila donde queda la firma del encuestador."""
    capacity = max(PERSONAS_PLANTILLA, person_count)
    extra = capacity - PERSONAS_PLANTILLA
    original_signature_row = 27

    if extra > 0:
        ws.insert_rows(original_signature_row, amount=extra)
        # La fila 26 tiene el formato normal de una persona y sirve de patrón.
        for row in range(original_signature_row, original_signature_row + extra):
            _copy_row_format(ws, 26, row, 30)  # A:AD

    return 12 + capacity


def _expand_food_sheet(ws, person_count: int) -> tuple[int, int]:
    """Amplía Hoja3 y devuelve (fila_firma, fila_pie)."""
    capacity = max(PERSONAS_PLANTILLA, person_count)
    extra_people = capacity - PERSONAS_PLANTILLA
    extra_rows = extra_people * 3
    original_signature_row = 57
    original_footer_row = 58

    # Esta combinación está debajo del área que se amplía. La recreamos al final.
    footer_merge = "B58:X58"
    if footer_merge in {str(rng) for rng in ws.merged_cells.ranges}:
        ws.unmerge_cells(footer_merge)

    if extra_rows > 0:
        ws.insert_rows(original_signature_row, amount=extra_rows)

        # Copiar el patrón de las tres filas de la persona 15 a cada nueva persona.
        for person_no in range(PERSONAS_PLANTILLA + 1, capacity + 1):
            base = 12 + (person_no - 1) * 3
            for offset in range(3):
                _copy_row_format(ws, 54 + offset, base + offset, 24)  # A:X
            ws.merge_cells(start_row=base, start_column=1, end_row=base + 2, end_column=1)
            ws.merge_cells(start_row=base, start_column=2, end_row=base + 2, end_column=2)

    signature_row = 12 + capacity * 3
    footer_row = signature_row + 1

    # Si no hubo inserción, el pie sigue en la fila 58. Si hubo, su contenido fue desplazado.
    ws.merge_cells(start_row=footer_row, start_column=2, end_row=footer_row, end_column=24)

    return signature_row, footer_row


def _write_general_dates(ws, general: dict) -> None:
    mappings = [
        (6, general.get("fecha_ocurrencia", "")),
        (7, general.get("fecha_notificacion", "")),
        (8, general.get("fecha_investigacion", "")),
    ]
    for row, text in mappings:
        d = _parse_date(text)
        if d:
            # Conserva las etiquetas originales y usa las celdas de espacio de la plantilla.
            ws[f"G{row}"] = d.day
            ws[f"I{row}"] = d.month
            ws[f"K{row}"] = d.year


def _write_symptom_headers(ws, symptoms: list[str]) -> None:
    for col in SINTOMA_COLS:
        ws[f"{col}10"] = ""
    for col, symptom in zip(SINTOMA_COLS, symptoms[:MAX_SINTOMAS]):
        cell = ws[f"{col}10"]
        cell.value = _clean_text(symptom)
        cell.alignment = copy(cell.alignment)
        cell.alignment = cell.alignment.copy(
            textRotation=90, wrapText=True, vertical="center", horizontal="center"
        )
    if symptoms:
        ws.row_dimensions[10].height = max(ws.row_dimensions[10].height or 15, 95)


def _write_people(ws, persons: list[dict], symptoms: list[str], capacity: int) -> None:
    symptom_map = {s: col for s, col in zip(symptoms[:MAX_SINTOMAS], SINTOMA_COLS)}
    for idx in range(capacity):
        row = 12 + idx
        p = persons[idx] if idx < len(persons) else {}
        ws[f"A{row}"] = idx + 1
        ws[f"B{row}"] = _clean_text(p.get("Nombres y apellidos"))
        ws[f"C{row}"] = _clean_text(p.get("Identificación"))
        age = p.get("Edad")
        ws[f"D{row}"] = None if _clean_text(age) == "" else age
        ws[f"E{row}"] = _clean_text(p.get("Sexo"))
        ws[f"F{row}"] = _clean_text(p.get("Dirección y teléfono"))
        ws[f"G{row}"] = _clean_text(p.get("Día síntomas"))
        ws[f"H{row}"] = _clean_text(p.get("Hora síntomas"))

        for col in SINTOMA_COLS:
            ws[f"{col}{row}"] = ""
        for symptom, col in symptom_map.items():
            if _truthy(p.get(symptom)):
                ws[f"{col}{row}"] = "X"

        ws[f"AA{row}"] = "Sí" if _truthy(p.get("Enfermo")) else ("No" if _clean_text(p.get("Enfermo")) != "" else "")
        ws[f"AB{row}"] = "Sí" if _truthy(p.get("Consulta")) else ("No" if _clean_text(p.get("Consulta")) != "" else "")
        ws[f"AC{row}"] = "Sí" if _truthy(p.get("Hospitalizado")) else ("No" if _clean_text(p.get("Hospitalizado")) != "" else "")
        ws[f"AD{row}"] = _clean_text(p.get("Muestra"))


def _write_food_headers(ws, foods: list[str]) -> None:
    for col in ALIMENTO_COLS:
        ws[f"{col}10"] = ""
    for col, food in zip(ALIMENTO_COLS, foods[:MAX_ALIMENTOS]):
        cell = ws[f"{col}10"]
        cell.value = _clean_text(food)
        cell.alignment = copy(cell.alignment)
        cell.alignment = cell.alignment.copy(
            textRotation=90, wrapText=True, vertical="center", horizontal="center"
        )
    if foods:
        ws.row_dimensions[10].height = max(ws.row_dimensions[10].height or 15, 95)


def _write_food_rows(ws, food_rows: list[dict], persons: list[dict], foods: list[str], capacity: int) -> None:
    food_map = {f: col for f, col in zip(foods[:MAX_ALIMENTOS], ALIMENTO_COLS)}
    indexed = {}
    for item in food_rows:
        try:
            no = int(float(item.get("No.")))
        except (TypeError, ValueError):
            continue
        period = _clean_text(item.get("Periodo"))
        indexed[(no, period)] = item

    for person_no in range(1, capacity + 1):
        person = persons[person_no - 1] if person_no - 1 < len(persons) else {}
        identification = _clean_text(person.get("Identificación"))
        base = 12 + (person_no - 1) * 3
        for offset, period in enumerate(PERIODOS):
            row = base + offset
            item = indexed.get((person_no, period), {})
            if offset == 0:
                ws[f"A{row}"] = person_no
                ws[f"B{row}"] = identification
            ws[f"C{row}"] = period
            ws[f"D{row}"] = _clean_text(item.get("Día"))
            ws[f"E{row}"] = _clean_text(item.get("Hora"))
            ws[f"F{row}"] = _clean_text(item.get("Lugar de consumo"))
            for col in ALIMENTO_COLS:
                ws[f"{col}{row}"] = ""
            for food, col in food_map.items():
                if _truthy(item.get(food)):
                    ws[f"{col}{row}"] = "X"


def build_excel(payload: dict, template_path: str | Path) -> tuple[bytes, str]:
    errors = validate_payload(payload)
    if errors:
        raise ValueError(" ".join(errors))

    template_path = Path(template_path)
    wb = load_workbook(template_path)
    ws1 = wb["Hoja1"]
    ws3 = wb["Hoja3"]

    general = payload.get("general", {})
    symptoms = [_clean_text(x) for x in payload.get("symptoms", []) if _clean_text(x)]
    foods = [_clean_text(x) for x in payload.get("foods", []) if _clean_text(x)]
    persons = payload.get("persons", [])[:MAX_PERSONAS]
    food_rows = payload.get("food_rows", [])

    person_count = max(1, min(len(persons), MAX_PERSONAS))
    capacity = max(PERSONAS_PLANTILLA, person_count)

    signature_row_1 = _expand_people_sheet(ws1, person_count)
    signature_row_3, _ = _expand_food_sheet(ws3, person_count)

    _write_general_dates(ws1, general)
    _write_symptom_headers(ws1, symptoms)
    _write_people(ws1, persons, symptoms, capacity)
    _write_food_headers(ws3, foods)
    _write_food_rows(ws3, food_rows, persons, foods, capacity)

    encuestador = _clean_text(general.get("encuestador"))
    telefono = _clean_text(general.get("telefono_encuestador"))
    ws1[f"B{signature_row_1}"] = (
        f"NOMBRE Y FIRMA DEL ENCUESTADOR: {encuestador}" if encuestador else "NOMBRE Y FIRMA DEL ENCUESTADOR"
    )
    ws1[f"H{signature_row_1}"] = f"Teléfono: {telefono}" if telefono else "Teléfono"
    ws3[f"B{signature_row_3}"] = (
        f"NOMBRE Y APELLIDOS - TELÉFONO Y FIRMA DEL ENCUESTADOR: {encuestador} - {telefono}".strip(" -")
        if (encuestador or telefono)
        else "NOMBRE Y APELLIDOS- TELEFONO Y FIRMA DEL ENCUESTADOR"
    )

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    date_part = _safe_filename(general.get("fecha_ocurrencia", ""))
    filename = f"ETA_Encuesta_{date_part}.xlsx"
    return output.getvalue(), filename
