from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE = "EAF2F8"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
THIN = Side(style="thin", color="808080")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AGE_GROUPS = [
    ("Menor de 1 año", None, 0),
    ("De 1 a 4 años", 1, 4),
    ("De 5 a 9 años", 5, 9),
    ("De 10 a 19 años", 10, 19),
    ("De 20 a 49 años", 20, 49),
    ("De 50 a 74 años", 50, 74),
    ("De 75 años y más", 75, None),
]


def _clean(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    return _clean(v).casefold() in {"sí", "si", "s", "1", "true", "x", "yes"}


def _fmt_date(v: Any) -> str:
    if not v:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%d/%m/%Y")
    text = _clean(v)
    if len(text) >= 10 and text[4] == "-":
        try:
            return datetime.strptime(text[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
        except ValueError:
            return text
    return text


def _fmt_time(v: Any) -> str:
    if not v:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    text = _clean(v)
    return text[:5] if len(text) >= 5 else text


def _safe_float(v: Any):
    try:
        if v in (None, ""):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def _active_person(p: dict) -> bool:
    return bool(_clean(p.get("Nombres y apellidos")) or _clean(p.get("Identificación")))


def _age_group(age) -> str:
    a = _safe_float(age)
    if a is None:
        return "Sin dato"
    for label, low, high in AGE_GROUPS:
        if low is None and a < 1:
            return label
        if high is None and a >= low:
            return label
        if low is not None and high is not None and low <= a <= high:
            return label
    return "Sin dato"


def analyze(payload: dict) -> dict:
    persons = [p for p in payload.get("persons", []) if _active_person(p)]
    foods = [f for f in payload.get("foods", []) if _clean(f)]
    symptoms = [s for s in payload.get("symptoms", []) if _clean(s)]
    food_rows = payload.get("food_rows", [])

    cases = [p for p in persons if _bool(p.get("Enfermo"))]
    noncases = [p for p in persons if not _bool(p.get("Enfermo"))]
    case_ids = {_clean(p.get("Identificación")) for p in cases if _clean(p.get("Identificación"))}
    total = len(persons)
    ncases = len(cases)

    symptom_counts = []
    for symptom in symptoms:
        count = sum(1 for p in cases if _bool(p.get(symptom)))
        pct = (count / ncases * 100) if ncases else 0
        symptom_counts.append((symptom, count, pct))
    symptom_counts.sort(key=lambda x: (-x[1], x[0].casefold()))

    sex = defaultdict(lambda: {"expuestos": 0, "casos": 0})
    age = defaultdict(lambda: {"expuestos": 0, "casos": 0})
    for p in persons:
        sx = _clean(p.get("Sexo")) or "Sin dato"
        sex[sx]["expuestos"] += 1
        grp = _age_group(p.get("Edad"))
        age[grp]["expuestos"] += 1
        if _bool(p.get("Enfermo")):
            sex[sx]["casos"] += 1
            age[grp]["casos"] += 1

    # Índice de consumo por No. de persona. Si consumió el alimento en cualquiera de los 3 periodos, cuenta como expuesto.
    consumed_by_no = defaultdict(set)
    for row in food_rows:
        try:
            no = int(float(row.get("No.")))
        except (TypeError, ValueError):
            continue
        for food in foods:
            if _bool(row.get(food)):
                consumed_by_no[no].add(food)

    active_by_no = {}
    for idx, p in enumerate(payload.get("persons", []), start=1):
        if _active_person(p):
            active_by_no[idx] = p

    food_analysis = []
    for food in foods:
        a = b = c = d = 0
        for no, p in active_by_no.items():
            sick = _bool(p.get("Enfermo"))
            exposed = food in consumed_by_no.get(no, set())
            if sick and exposed:
                a += 1
            elif (not sick) and exposed:
                b += 1
            elif sick and (not exposed):
                c += 1
            else:
                d += 1
        ar_exp = a / (a + b) if (a + b) else None
        ar_unexp = c / (c + d) if (c + d) else None
        corrected = any(v == 0 for v in (a, b, c, d))
        aa, bb, cc, dd = ((a + 1, b + 1, c + 1, d + 1) if corrected else (a, b, c, d))
        ar_exp_c = aa / (aa + bb) if (aa + bb) else None
        ar_unexp_c = cc / (cc + dd) if (cc + dd) else None
        rr = (ar_exp_c / ar_unexp_c) if (ar_unexp_c not in (None, 0)) else None
        orr = (aa * dd / (bb * cc)) if bb and cc else None
        food_analysis.append({
            "food": food, "a": a, "b": b, "c": c, "d": d,
            "attack_exposed": ar_exp, "attack_unexposed": ar_unexp,
            "rr": rr, "or": orr, "corrected": corrected,
        })

    onset_counts = Counter()
    for p in cases:
        d = _fmt_date(p.get("Día síntomas"))
        if d:
            onset_counts[d] += 1

    return {
        "persons": persons,
        "cases": cases,
        "noncases": noncases,
        "total_exposed": total,
        "total_cases": ncases,
        "attack_rate": (ncases / total * 100) if total else 0,
        "hospitalized": sum(1 for p in cases if _bool(p.get("Hospitalizado"))),
        "consulted": sum(1 for p in cases if _bool(p.get("Consulta"))),
        "with_sample": sum(1 for p in cases if _clean(p.get("Muestra"))),
        "symptom_counts": symptom_counts,
        "sex": dict(sex),
        "age": dict(age),
        "food_analysis": food_analysis,
        "onset_counts": dict(onset_counts),
    }


def _setup_sheet(ws, title: str):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H2")
    c = ws["A1"]
    c.value = title
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.font = Font(color=WHITE, bold=True, size=14)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, width in enumerate([23, 23, 18, 18, 18, 18, 18, 22], 1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _section(ws, row: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    c = ws.cell(row, 1, text)
    c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    c.font = Font(bold=True, color=BLUE)
    c.alignment = Alignment(vertical="center")
    c.border = BORDER
    return row + 1


def _kv(ws, row: int, pairs: list[tuple[str, Any]]) -> int:
    # up to 2 pairs per row, each pair spans label/value across 4 cols
    for i, (label, value) in enumerate(pairs[:2]):
        start = 1 + i * 4
        ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=start + 1)
        ws.merge_cells(start_row=row, start_column=start + 2, end_row=row, end_column=start + 3)
        lc = ws.cell(row, start, label)
        vc = ws.cell(row, start + 2, value if value is not None else "")
        lc.fill = PatternFill("solid", fgColor=GRAY)
        lc.font = Font(bold=True)
        lc.alignment = Alignment(wrap_text=True, vertical="center")
        vc.alignment = Alignment(wrap_text=True, vertical="center")
        for c in range(start, start + 4):
            ws.cell(row, c).border = BORDER
    ws.row_dimensions[row].height = 32
    return row + 1


def _text_block(ws, row: int, label: str, text: Any, height: int = 60) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=8)
    lc = ws.cell(row, 1, label)
    lc.fill = PatternFill("solid", fgColor=GRAY)
    lc.font = Font(bold=True)
    lc.alignment = Alignment(wrap_text=True, vertical="top")
    vc = ws.cell(row, 3, _clean(text))
    vc.alignment = Alignment(wrap_text=True, vertical="top")
    for rr in range(row, row + 2):
        for cc in range(1, 9):
            ws.cell(rr, cc).border = BORDER
    ws.row_dimensions[row].height = height / 2
    ws.row_dimensions[row + 1].height = height / 2
    return row + 2


def _table(ws, row: int, headers: list[str], rows: list[list[Any]], widths=None) -> int:
    n = len(headers)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col, h)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 34
    row += 1
    for data in rows:
        for col in range(1, n + 1):
            val = data[col - 1] if col - 1 < len(data) else ""
            c = ws.cell(row, col, val)
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
        row += 1
    return row


def _auto_summary(analysis: dict) -> str:
    return (
        f"Se consolidaron {analysis['total_exposed']} personas expuestas, de las cuales "
        f"{analysis['total_cases']} fueron clasificadas como enfermas. La tasa de ataque general "
        f"es {analysis['attack_rate']:.1f}%. Se registran {analysis['hospitalized']} hospitalizados "
        f"y {analysis['consulted']} personas que consultaron a servicios de salud."
    )


def build_reports(payload: dict, report_fields: dict | None = None) -> tuple[bytes, str]:
    report_fields = report_fields or {}
    a = analyze(payload)
    g = payload.get("general", {})

    wb = Workbook()
    ws24 = wb.active
    ws24.title = "Inf. Preliminar"
    ws72 = wb.create_sheet("Inf. 72 horas")
    wsf = wb.create_sheet("Inf. final")

    # 24 horas
    _setup_sheet(ws24, "INFORME PRELIMINAR - INVESTIGACIÓN DE POSIBLE BROTE DE ETA")
    r = 4
    r = _section(ws24, r, "1. Identificación del brote")
    r = _kv(ws24, r, [("Departamento", g.get("departamento", "")), ("Municipio", g.get("municipio", ""))])
    r = _kv(ws24, r, [("Barrio/Corregimiento/Vereda", g.get("localidad", "")), ("Lugar / institución", g.get("lugar_brote", ""))])
    r = _kv(ws24, r, [("Dirección", g.get("direccion_brote", "")), ("Teléfono", g.get("telefono_brote", ""))])
    r = _kv(ws24, r, [("Fecha inicio del brote", _fmt_date(g.get("fecha_ocurrencia"))), ("Fecha detección", _fmt_date(g.get("fecha_deteccion")))])
    r = _kv(ws24, r, [("Fecha notificación", _fmt_date(g.get("fecha_notificacion"))), ("Fecha inicio IEC", _fmt_date(g.get("fecha_investigacion")))])
    r = _section(ws24, r, "2. Situación inicial")
    r = _kv(ws24, r, [("Posible número de expuestos", a["total_exposed"]), ("Total de casos", a["total_cases"])])
    r = _kv(ws24, r, [("Casos UPGD", report_fields.get("casos_upgd", "")), ("Casos BAC", report_fields.get("casos_bac", ""))])
    r = _kv(ws24, r, [("Casos hospitalizados", a["hospitalized"]), ("Tasa de ataque", f"{a['attack_rate']:.1f}%")])
    r = _text_block(ws24, r, "Antecedentes del brote", report_fields.get("antecedentes", _auto_summary(a)))
    top_sym = ", ".join(f"{s} ({n})" for s, n, _ in a["symptom_counts"][:8])
    r = _text_block(ws24, r, "Signos y síntomas principales", top_sym)
    foods_named = ", ".join(payload.get("foods", []))
    r = _text_block(ws24, r, "Posibles alimentos/agua o mecanismos de transmisión", report_fields.get("posibles_alimentos", foods_named))
    r = _text_block(ws24, r, "Hipótesis inicial", report_fields.get("hipotesis_inicial", ""))
    r = _text_block(ws24, r, "Medidas iniciales de control", report_fields.get("medidas_control", ""))
    r = _text_block(ws24, r, "Otra información relevante", report_fields.get("otra_informacion", ""))
    r = _kv(ws24, r, [("Responsable", report_fields.get("responsable", g.get("encuestador", ""))), ("Teléfono", g.get("telefono_encuestador", ""))])
    ws24.freeze_panes = "A4"

    # 72 horas
    _setup_sheet(ws72, "SEGUNDO INFORME - 72 HORAS - INVESTIGACIÓN DE BROTE ETA")
    r = 4
    r = _section(ws72, r, "1. Datos generales y descripción")
    r = _kv(ws72, r, [("Departamento", g.get("departamento", "")), ("Municipio", g.get("municipio", ""))])
    r = _kv(ws72, r, [("Total expuestos", a["total_exposed"]), ("Total casos", a["total_cases"])])
    r = _kv(ws72, r, [("Hospitalizados", a["hospitalized"]), ("Con muestra", a["with_sample"])])
    r = _text_block(ws72, r, "Antecedentes del brote", report_fields.get("antecedentes", _auto_summary(a)))
    r = _text_block(ws72, r, "Definición operacional de caso", report_fields.get("definicion_caso", ""))
    r = _text_block(ws72, r, "Manejo y tratamiento clínico / complicaciones", report_fields.get("manejo_clinico", ""))

    r = _section(ws72, r, "2. Distribución por edad y sexo")
    age_rows = []
    for label, _, _ in AGE_GROUPS:
        d = a["age"].get(label, {"expuestos": 0, "casos": 0})
        age_rows.append([label, d["expuestos"], d["casos"], f"{(d['casos']/d['expuestos']*100 if d['expuestos'] else 0):.1f}%"])
    age_rows.append(["TOTAL", a["total_exposed"], a["total_cases"], f"{a['attack_rate']:.1f}%"])
    r = _table(ws72, r, ["Grupo de edad", "Expuestos", "Casos", "Tasa de ataque"], age_rows)
    r += 1
    sex_rows = [[k, v["expuestos"], v["casos"], f"{(v['casos']/v['expuestos']*100 if v['expuestos'] else 0):.1f}%"] for k, v in sorted(a["sex"].items())]
    r = _table(ws72, r, ["Sexo", "Expuestos", "Casos", "Tasa de ataque"], sex_rows)

    r += 1
    r = _section(ws72, r, "3. Distribución de casos por signos y síntomas")
    sym_rows = [[s, n, f"{pct:.1f}%"] for s, n, pct in a["symptom_counts"]]
    r = _table(ws72, r, ["Signo / síntoma", "No. de casos", "Porcentaje"], sym_rows or [["Sin datos", 0, "0.0%"]])

    r += 1
    r = _section(ws72, r, "4. Análisis epidemiológico por alimento")
    fa_rows = []
    for x in a["food_analysis"]:
        fa_rows.append([
            x["food"], x["a"], x["b"], x["c"], x["d"],
            f"{x['attack_exposed']*100:.1f}%" if x["attack_exposed"] is not None else "N/A",
            f"{x['attack_unexposed']*100:.1f}%" if x["attack_unexposed"] is not None else "N/A",
            f"{x['rr']:.2f}" if x["rr"] is not None else "N/A",
        ])
    # 8 cols max; split OR into next concise table if needed.
    r = _table(ws72, r, ["Alimento", "Caso exp.", "Sano exp.", "Caso no exp.", "Sano no exp.", "TA exp.", "TA no exp.", "RR*"], fa_rows or [["Sin alimentos definidos", 0, 0, 0, 0, "N/A", "N/A", "N/A"]])
    ws72.cell(r, 1, "*Si alguna celda de la tabla 2x2 es cero, el cálculo de RR/OR aplica la corrección indicada en el formato original (suma 1 a las cuatro celdas).")
    ws72.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws72.cell(r, 1).font = Font(italic=True, size=9)
    ws72.cell(r, 1).alignment = Alignment(wrap_text=True)
    r += 2

    r = _section(ws72, r, "5. Curva epidémica")
    curve_start = r
    curve_rows = [[d, n] for d, n in sorted(a["onset_counts"].items())]
    r = _table(ws72, r, ["Fecha de inicio de síntomas", "No. de casos"], curve_rows or [["Sin fecha registrada", 0]])
    if curve_rows:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Curva epidémica"
        chart.y_axis.title = "Número de casos"
        chart.x_axis.title = "Fecha"
        data = Reference(ws72, min_col=2, min_row=curve_start, max_row=curve_start + len(curve_rows))
        cats = Reference(ws72, min_col=1, min_row=curve_start + 1, max_row=curve_start + len(curve_rows))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 7
        chart.width = 14
        ws72.add_chart(chart, f"D{curve_start}")
        r = max(r, curve_start + 14)

    r = _section(ws72, r, "6. Hallazgos, laboratorio y control")
    r = _text_block(ws72, r, "Resultados de laboratorio", report_fields.get("resultados_laboratorio", ""))
    r = _text_block(ws72, r, "Hallazgos ambientales / factores de riesgo", report_fields.get("hallazgos_ambientales", ""))
    r = _text_block(ws72, r, "Análisis de resultados e hipótesis", report_fields.get("analisis_resultados", ""))
    r = _text_block(ws72, r, "Medidas de control", report_fields.get("medidas_control_72", report_fields.get("medidas_control", "")))
    r = _text_block(ws72, r, "Recomendaciones", report_fields.get("recomendaciones_72", ""))
    r = _text_block(ws72, r, "Conclusiones", report_fields.get("conclusiones_72", ""))
    r = _kv(ws72, r, [("Responsable", report_fields.get("responsable", g.get("encuestador", ""))), ("Teléfono", g.get("telefono_encuestador", ""))])
    ws72.freeze_panes = "A4"

    # Final
    _setup_sheet(wsf, "INFORME FINAL - INVESTIGACIÓN DE BROTE ETA")
    r = 4
    r = _section(wsf, r, "1. Resumen de la situación")
    r = _kv(wsf, r, [("Departamento", g.get("departamento", "")), ("Municipio", g.get("municipio", ""))])
    r = _kv(wsf, r, [("Total expuestos", a["total_exposed"]), ("Total casos", a["total_cases"])])
    r = _kv(wsf, r, [("Tasa de ataque", f"{a['attack_rate']:.1f}%"), ("Hospitalizados", a["hospitalized"])])
    r = _kv(wsf, r, [("Fecha inicio", _fmt_date(g.get("fecha_ocurrencia"))), ("Fecha cierre", _fmt_date(report_fields.get("fecha_cierre")))])
    r = _text_block(wsf, r, "Resumen de la situación", report_fields.get("resumen_final", _auto_summary(a)))
    r = _text_block(wsf, r, "Descripción del brote", report_fields.get("descripcion_brote", report_fields.get("antecedentes", "")))
    r = _section(wsf, r, "2. Caracterización y cierre")
    r = _kv(wsf, r, [("Estado del brote", report_fields.get("estado_brote", "")), ("Agente identificado", report_fields.get("agente_identificado", ""))])
    r = _kv(wsf, r, [("Fuente / alimento implicado", report_fields.get("fuente_implicada", "")), ("Modo de transmisión", report_fields.get("modo_transmision", ""))])
    r = _text_block(wsf, r, "Resultados de laboratorio", report_fields.get("resultados_laboratorio_final", report_fields.get("resultados_laboratorio", "")))
    r = _text_block(wsf, r, "Factores determinantes", report_fields.get("factores_determinantes", report_fields.get("hallazgos_ambientales", "")))
    r = _section(wsf, r, "3. Síntomas y análisis por alimento")
    sym_rows = [[s, n, f"{pct:.1f}%"] for s, n, pct in a["symptom_counts"]]
    r = _table(wsf, r, ["Signo / síntoma", "No. casos", "Porcentaje"], sym_rows or [["Sin datos", 0, "0.0%"]])
    r += 1
    fa_rows2 = []
    for x in a["food_analysis"]:
        fa_rows2.append([
            x["food"], x["a"], x["b"], x["c"], x["d"],
            f"{x['rr']:.2f}" if x["rr"] is not None else "N/A",
            f"{x['or']:.2f}" if x["or"] is not None else "N/A",
            "Sí" if x["corrected"] else "No",
        ])
    r = _table(wsf, r, ["Alimento", "Caso exp.", "Sano exp.", "Caso no exp.", "Sano no exp.", "RR", "OR", "Corrección"], fa_rows2 or [["Sin alimentos", 0,0,0,0,"N/A","N/A","No"]])
    r += 1
    r = _section(wsf, r, "4. Recomendaciones, conclusiones y seguimiento")
    r = _text_block(wsf, r, "Recomendaciones", report_fields.get("recomendaciones_final", report_fields.get("recomendaciones_72", "")))
    r = _text_block(wsf, r, "Conclusiones", report_fields.get("conclusiones_final", report_fields.get("conclusiones_72", "")))
    r = _text_block(wsf, r, "Plan de mejoramiento / seguimiento", report_fields.get("plan_mejoramiento", ""), height=80)
    r = _kv(wsf, r, [("Responsable", report_fields.get("responsable", g.get("encuestador", ""))), ("Teléfono", g.get("telefono_encuestador", ""))])
    wsf.freeze_panes = "A4"

    for ws in (ws24, ws72, wsf):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    dt = _fmt_date(g.get("fecha_ocurrencia")).replace("/", "-") or "sin_fecha"
    return out.getvalue(), f"ETA_Informes_24_72_Final_{dt}.xlsx"
